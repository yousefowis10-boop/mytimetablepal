"""Excel / CSV import (with dry-run preview) and Excel exports."""
import csv
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from api.access import SchoolAccess, log
from api.models import Classroom, StudentGroup, Subject, Teacher, Timetable

from ..common import MANAGERS, active_timetable, parse_date
from ..models import AttendanceRecord, Grade, Guardian, Student, StudentGuardian, Term

COLUMNS = {
    "students": ["first_name", "last_name", "student_number", "group", "date_of_birth", "gender",
                 "parent_name", "parent_phone", "parent_email", "parent_relation",
                 "second_parent_name", "second_parent_phone", "emergency_contact_name", "emergency_contact_phone",
                 "medical_notes", "allergies", "home_address", "uses_bus_morning", "uses_bus_afternoon"],
    "teachers": ["name", "subjects", "phone", "email", "experience", "max_lessons_per_day", "max_lessons_per_week"],
    "subjects": ["name", "code", "requires_room_type", "is_heavy"],
    "classrooms": ["name", "capacity", "room_type"],
    "groups": ["name", "grade_level", "size", "home_room"],
}
REQUIRED = {"students": ["first_name"], "teachers": ["name"], "subjects": ["name", "code"],
            "classrooms": ["name", "capacity"], "groups": ["name"]}
EXAMPLE = {
    "students": ["Lina", "Haddad", "S-1001", "Grade 3-A", "2017-04-12", "F", "Omar Haddad", "0791234567",
                 "omar@example.com", "Father", "Rania Haddad", "0797654321", "Rania Haddad", "0797654321",
                 "", "Peanuts", "Khalda, Amman", "yes", "yes"],
    "teachers": ["Sara Odeh", "Mathematics, Physics", "0790000000", "sara@example.com", "8", "6", "24"],
    "subjects": ["Mathematics", "MATH", "", "yes"],
    "classrooms": ["Room 101", "32", "regular"],
    "groups": ["Grade 3-A", "3", "28", "Room 101"],
}


def truthy(v):
    return str(v).strip().lower() in ("1", "yes", "y", "true", "نعم", "x")


def read_rows(upload):
    name = upload.name.lower()
    if name.endswith(".csv"):
        text = upload.read().decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    elif name.endswith((".xlsx", ".xlsm")):
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        raise ValidationError({"file": "Upload an .xlsx or .csv file."})
    if not rows:
        raise ValidationError({"file": "The file is empty."})
    header = [str(h).strip().lower().replace(" ", "_") for h in rows[0]]
    out = []
    for i, r in enumerate(rows[1:], start=2):
        if not any(str(c).strip() for c in r):
            continue
        out.append((i, {header[j]: (r[j] if j < len(r) else "") for j in range(len(header))}))
    if len(out) > 5000:
        raise ValidationError({"file": "Please import at most 5,000 rows at a time."})
    return out


def cell(row, key):
    v = row.get(key, "")
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


class ImportView(APIView):
    """GET  /schools/{id}/import/{kind}/  -> Excel template
       POST /schools/{id}/import/{kind}/  file=..., commit=0|1 -> preview (default) or import"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def get(self, request, school_id, kind):
        SchoolAccess(request, school_id, read_roles=MANAGERS)
        if kind not in COLUMNS:
            raise ValidationError("Unknown import type.")
        wb = Workbook()
        ws = wb.active
        ws.title = kind
        ws.append(COLUMNS[kind])
        ws.append(EXAMPLE[kind])
        for c in ws[1]:
            c.font = Font(bold=True)
        return xlsx_response(wb, f"{kind}-template.xlsx")

    def post(self, request, school_id, kind):
        access = SchoolAccess(request, school_id, read_roles=MANAGERS, write_roles=MANAGERS)
        if kind not in COLUMNS:
            raise ValidationError("Unknown import type.")
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": "Choose a file."})
        rows = read_rows(upload)
        commit = truthy(request.data.get("commit", "0"))
        importer = getattr(self, f"import_{kind}")
        errors, preview = [], []
        try:
            with transaction.atomic():
                created, updated = importer(access.school, rows, errors, preview)
                if errors or not commit:
                    raise _Rollback()
        except _Rollback:
            if errors or not commit:
                return Response({"committed": False, "rows": len(rows), "errors": errors[:200],
                                 "preview": preview[:20], "would_create": created, "would_update": updated})
        log(access.school, request.user, "create", None, model=f"import.{kind}",
            summary=f"Imported {kind}: {created} new, {updated} updated")
        return Response({"committed": True, "rows": len(rows), "created": created, "updated": updated, "errors": []})

    # --- importers: return (created, updated); append (line, message) to errors ---
    def import_subjects(self, school, rows, errors, preview):
        created = updated = 0
        for line, r in rows:
            name, code = cell(r, "name"), cell(r, "code").upper()
            if not name or not code:
                errors.append({"line": line, "error": "name and code are required"})
                continue
            rtype = cell(r, "requires_room_type")
            if rtype and rtype not in ("lab", "computer_lab", "gym", "other", "regular"):
                errors.append({"line": line, "error": f"unknown room type '{rtype}'"})
                continue
            _, new = Subject.objects.update_or_create(school=school, code=code, defaults={
                "name": name, "requires_room_type": "" if rtype == "regular" else rtype, "is_heavy": truthy(cell(r, "is_heavy"))})
            created += new
            updated += not new
            preview.append({"line": line, "name": name, "code": code})
        return created, updated

    def import_classrooms(self, school, rows, errors, preview):
        created = updated = 0
        for line, r in rows:
            name = cell(r, "name")
            try:
                cap = int(cell(r, "capacity"))
            except ValueError:
                errors.append({"line": line, "error": "capacity must be a number"})
                continue
            rtype = cell(r, "room_type") or "regular"
            if rtype not in ("regular", "lab", "computer_lab", "gym", "other"):
                errors.append({"line": line, "error": f"unknown room type '{rtype}'"})
                continue
            obj = Classroom.objects.filter(school=school, name__iexact=name).first()
            if obj:
                obj.capacity, obj.room_type = cap, rtype
                obj.save()
                updated += 1
            else:
                Classroom.objects.create(school=school, name=name, capacity=cap, room_type=rtype)
                created += 1
            preview.append({"line": line, "name": name, "capacity": cap})
        return created, updated

    def import_groups(self, school, rows, errors, preview):
        created = updated = 0
        rooms = {c.name.lower(): c for c in Classroom.objects.filter(school=school)}
        for line, r in rows:
            name = cell(r, "name")
            room = rooms.get(cell(r, "home_room").lower()) if cell(r, "home_room") else None
            if cell(r, "home_room") and not room:
                errors.append({"line": line, "error": f"room '{cell(r, 'home_room')}' not found - import rooms first"})
                continue
            defaults = {"classroom": room}
            for f in ("grade_level", "size"):
                if cell(r, f):
                    try:
                        defaults[f] = int(cell(r, f))
                    except ValueError:
                        errors.append({"line": line, "error": f"{f} must be a number"})
            obj = StudentGroup.objects.filter(school=school, name__iexact=name).first()
            if obj:
                for k, v in defaults.items():
                    setattr(obj, k, v)
                obj.save()
                updated += 1
            else:
                StudentGroup.objects.create(school=school, name=name, **defaults)
                created += 1
            preview.append({"line": line, "name": name})
        return created, updated

    def import_teachers(self, school, rows, errors, preview):
        created = updated = 0
        subjects = {s.name.lower(): s for s in Subject.objects.filter(school=school)}
        subjects.update({s.code.lower(): s for s in subjects.values()})
        for line, r in rows:
            name = cell(r, "name")
            subj_names = [x.strip().lower() for x in cell(r, "subjects").replace("،", ",").split(",") if x.strip()]
            missing = [x for x in subj_names if x not in subjects]
            if missing:
                errors.append({"line": line, "error": f"unknown subject(s): {', '.join(missing)}"})
                continue
            t = Teacher.objects.filter(school=school, name__iexact=name).first()
            new = t is None
            if new:
                t = Teacher(school=school, name=name, experience=0)
            for f in ("phone", "email"):
                if cell(r, f):
                    setattr(t, f, cell(r, f))
            for f in ("experience", "max_lessons_per_day", "max_lessons_per_week"):
                if cell(r, f):
                    try:
                        setattr(t, f, int(cell(r, f)))
                    except ValueError:
                        errors.append({"line": line, "error": f"{f} must be a number"})
            t.save()
            t.subjects.set([subjects[x] for x in subj_names])
            created += new
            updated += not new
            preview.append({"line": line, "name": name, "subjects": subj_names})
        return created, updated

    def import_students(self, school, rows, errors, preview):
        created = updated = 0
        groups = {g.name.lower(): g for g in StudentGroup.objects.filter(school=school)}
        active = Student.objects.filter(school=school, is_active=True).count()
        for line, r in rows:
            first = cell(r, "first_name")
            if not first:
                errors.append({"line": line, "error": "first_name is required"})
                continue
            group = groups.get(cell(r, "group").lower()) if cell(r, "group") else None
            if cell(r, "group") and not group:
                errors.append({"line": line, "error": f"class group '{cell(r, 'group')}' not found"})
                continue
            dob = None
            if cell(r, "date_of_birth"):
                try:
                    dob = date.fromisoformat(cell(r, "date_of_birth")[:10])
                except ValueError:
                    errors.append({"line": line, "error": "date_of_birth must be YYYY-MM-DD"})
                    continue
            number = cell(r, "student_number")
            s = Student.objects.filter(school=school, student_number=number).first() if number else None
            new = s is None
            if new:
                if active >= school.max_students:
                    errors.append({"line": line, "error": f"plan limit of {school.max_students} students reached"})
                    continue
                s = Student(school=school)
                active += 1
            s.first_name, s.last_name, s.student_number, s.group = first, cell(r, "last_name"), number, group
            s.date_of_birth = dob
            s.gender = cell(r, "gender")[:1].upper() if cell(r, "gender")[:1].upper() in ("M", "F") else ""
            for f in ("emergency_contact_name", "emergency_contact_phone", "medical_notes", "allergies", "home_address"):
                setattr(s, f, cell(r, f))
            s.uses_bus_morning = truthy(cell(r, "uses_bus_morning"))
            s.uses_bus_afternoon = truthy(cell(r, "uses_bus_afternoon"))
            s.save()
            for prefix, primary in (("parent", True), ("second_parent", False)):
                pname = cell(r, f"{prefix}_name")
                if not pname:
                    continue
                phone = cell(r, f"{prefix}_phone")
                g = (Guardian.objects.filter(school=school, phone=phone).first() if phone else None) or \
                    Guardian.objects.filter(school=school, name__iexact=pname).first()
                if not g:
                    g = Guardian.objects.create(school=school, name=pname, phone=phone,
                                                email=cell(r, f"{prefix}_email"), relation=cell(r, f"{prefix}_relation"))
                StudentGuardian.objects.get_or_create(student=s, guardian=g, defaults={"is_primary": primary})
            created += new
            updated += not new
            preview.append({"line": line, "name": f"{first} {s.last_name}".strip(),
                            "group": group.name if group else ""})
        return created, updated


class _Rollback(Exception):
    pass


def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(),
                        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def styled(ws, header):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")


class ExportView(APIView):
    """GET /schools/{id}/export/{kind}/ -> .xlsx
    kinds: attendance (?from=&to=), grades (?term=), fees, timetable (?type=grade|teacher|room)"""
    permission_classes = [IsAuthenticated]

    def get(self, request, school_id, kind):
        access = SchoolAccess(request, school_id, read_roles=("admin", "staff", "teacher"))
        school = access.school
        wb = Workbook()
        ws = wb.active
        p = request.query_params
        if kind == "attendance":
            end = parse_date(p.get("to"), timezone.localdate())
            start = parse_date(p.get("from"), end - timedelta(days=30))
            styled(ws, ["Date", "Student", "Group", "Period", "Status", "Note"])
            for r in AttendanceRecord.objects.filter(school=school, date__gte=start, date__lte=end) \
                    .select_related("student__group", "period").order_by("date", "student__first_name"):
                ws.append([r.date, r.student.full_name, r.student.group.name if r.student.group else "",
                           r.period.name if r.period else "", r.status, r.note])
        elif kind == "grades":
            term = Term.objects.filter(school=school, id=p.get("term")).first() if p.get("term") else \
                Term.objects.filter(school=school).order_by("-start_date").first()
            styled(ws, ["Term", "Group", "Subject", "Assessment", "Date", "Student", "Score", "Max", "%"])
            for g in Grade.objects.filter(assessment__school=school, assessment__term=term) \
                    .select_related("assessment__group", "assessment__subject", "student"):
                a = g.assessment
                pct = round(float(g.score / a.max_score * 100), 1) if g.score is not None and a.max_score else None
                ws.append([term.name if term else "", a.group.name, a.subject.name, a.name, a.date,
                           g.student.full_name, float(g.score) if g.score is not None else None, float(a.max_score), pct])
        elif kind == "fees":
            if not access.is_manager:
                raise ValidationError("Only admin and staff export fees.")
            from ..models import Payment, StudentFee
            styled(ws, ["Student", "Group", f"Billed ({school.currency})", "Paid", "Balance"])
            billed, paid = defaultdict(float), defaultdict(float)
            for f in StudentFee.objects.filter(school=school):
                billed[f.student_id] += float(f.amount - f.discount)
            for pmt in Payment.objects.filter(school=school):
                paid[pmt.student_id] += float(pmt.amount)
            for s in Student.objects.filter(school=school).select_related("group"):
                if s.id in billed or s.id in paid:
                    ws.append([s.full_name, s.group.name if s.group else "", round(billed[s.id], 3),
                               round(paid[s.id], 3), round(billed[s.id] - paid[s.id], 3)])
        elif kind == "timetable":
            from api.views import build_timetable_pages
            tt = Timetable.objects.filter(school=school, id=p.get("timetable")).first() if p.get("timetable") \
                else active_timetable(school)
            if not tt:
                raise ValidationError("No timetable yet.")
            data = build_timetable_pages(tt, p.get("type", "grade"))
            from api.scheduling.solver import DAY_NAMES
            wb.remove(ws)
            for page in data["pages"]:
                sheet = wb.create_sheet(page["title"].split(": ", 1)[-1][:31])
                styled(sheet, ["Period"] + [DAY_NAMES[d] for d in data["days"]])
                for row in page["periods"]:
                    sheet.append([row["period"]] + [
                        " / ".join(x for x in (c.get("subject"), c.get("teacher") or c.get("student_group"),
                                               c.get("classroom") or c.get("student_group")) if x) if c["subject"] else ""
                        for c in row["days"]])
            if not wb.worksheets:
                wb.create_sheet("Empty")
        else:
            raise ValidationError("Unknown export.")
        log(school, request.user, "view", None, model=f"export.{kind}", summary=f"Exported {kind}")
        return xlsx_response(wb, f"{school.name}-{kind}-{timezone.localdate()}.xlsx".replace(" ", "_"))
